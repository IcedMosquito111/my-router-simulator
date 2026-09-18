# -*- coding: utf-8 -*-
"""
生成《计算机与网络课程设计 - 课程学习体会》DOCX。

输入：计算机与网络课程设计-课程学习体会-模板.docx（学校模板：页眉标题 + 学号/姓名/签字/日期表）
      report/reflection_source.md（体会正文源）
输出：report/课程学习体会-王文聪-1120241345.docx

用法：
    python report/scripts/build_reflection_docx.py
"""

import pathlib
import sys

from docx import Document
from docx.shared import Pt

ROOT = pathlib.Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "计算机与网络课程设计-课程学习体会-模板.docx"
SOURCE = ROOT / "report" / "reflection_source.md"
OUTPUT = ROOT / "report" / "课程学习体会-王文聪-1120241345.docx"

BODY_FONT_SIZE = Pt(12)


def parse_source(text: str):
    """解析体会源：元信息字典 + 正文段落列表（支持 **加粗** 标记）"""
    meta, body = {}, []
    in_body = False
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#") or not stripped:
            continue
        if stripped.startswith("正文："):
            in_body = True
            continue
        if not in_body and "：" in stripped:
            key, value = stripped.split("：", 1)
            meta[key.strip()] = value.strip()
            continue
        body.append(stripped)
    return meta, body


def add_rich_paragraph(document: Document, text: str) -> None:
    """写入一段正文，处理 **加粗** 与首行缩进"""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(24)   # 首行缩进 2 字符
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    for index, chunk in enumerate(text.split("**")):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        run.font.size = BODY_FONT_SIZE
        run.bold = index % 2 == 1     # 奇数段落在 ** 之间，即加粗内容


def fill_header_table(document: Document, meta: dict) -> None:
    cells = document.tables[0].rows[0].cells
    cells[1].text = meta.get("学号", "")
    cells[3].text = meta.get("姓名", "")
    cells[5].text = ""                       # 签字留空由本人手写
    cells[7].text = meta.get("日期", "")
    for cell in cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(10.5)


def main() -> None:
    if not TEMPLATE.is_file():
        sys.exit(f"未找到体会模板：{TEMPLATE}")
    meta, body = parse_source(SOURCE.read_text(encoding="utf-8-sig"))

    document = Document(str(TEMPLATE))
    fill_header_table(document, meta)

    # 模板正文只有两个空段落，清掉后按需写入
    for paragraph in list(document.paragraphs):
        if not paragraph.text.strip():
            paragraph._element.getparent().remove(paragraph._element)
    for text in body:
        add_rich_paragraph(document, text)

    document.save(str(OUTPUT))
    characters = sum(len(t.replace("**", "")) for t in body)
    print(f"学习体会已生成：{OUTPUT}")
    print(f"  正文 {len(body)} 段，约 {characters} 字")


if __name__ == "__main__":
    main()
